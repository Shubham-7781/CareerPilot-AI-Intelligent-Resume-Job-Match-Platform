import io
import docx


def _signup_and_login(client, email="resumeuser@example.com"):
    client.post("/api/auth/signup", json={
        "full_name": "Resume User", "email": email, "password": "strongpassword123",
    })
    resp = client.post("/api/auth/login", data={"username": email, "password": "strongpassword123"})
    token = resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _make_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("jane.doe@example.com")
    document.add_paragraph("Skills")
    document.add_paragraph("Python, React, PostgreSQL, Docker")
    document.add_paragraph("Experience")
    document.add_paragraph("Backend Engineer at Example Corp — built REST APIs with FastAPI.")
    document.add_paragraph("Education")
    document.add_paragraph("B.Tech Computer Science, 2024")
    document.add_paragraph("Projects")
    document.add_paragraph("CareerPilot — resume analysis platform.")
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    return buf.read()


def test_upload_requires_auth(client):
    resp = client.post("/api/resumes/upload", files={"file": ("resume.docx", b"fake", "application/octet-stream")})
    assert resp.status_code == 401


def test_upload_rejects_disallowed_extension(client):
    headers = _signup_and_login(client)
    resp = client.post(
        "/api/resumes/upload",
        headers=headers,
        files={"file": ("resume.exe", b"MZfakebinary", "application/octet-stream")},
    )
    assert resp.status_code == 400


def test_upload_rejects_content_that_does_not_match_extension(client):
    headers = _signup_and_login(client)
    # .pdf extension but the content isn't a real PDF (no %PDF- magic bytes)
    resp = client.post(
        "/api/resumes/upload",
        headers=headers,
        files={"file": ("resume.pdf", b"this is not a real pdf", "application/pdf")},
    )
    assert resp.status_code == 400


def test_upload_valid_docx_and_parses_skills(client):
    headers = _signup_and_login(client)
    docx_bytes = _make_docx_bytes()

    resp = client.post(
        "/api/resumes/upload",
        headers=headers,
        files={"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 201
    data = resp.json()
    assert data["original_filename"] == "resume.docx"
    assert data["parsed_sections"] is not None


def test_list_and_get_and_delete_resume(client):
    headers = _signup_and_login(client, email="crud@example.com")
    docx_bytes = _make_docx_bytes()

    upload_resp = client.post(
        "/api/resumes/upload",
        headers=headers,
        files={"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    resume_id = upload_resp.json()["id"]

    list_resp = client.get("/api/resumes/", headers=headers)
    assert list_resp.status_code == 200
    assert any(r["id"] == resume_id for r in list_resp.json())

    get_resp = client.get(f"/api/resumes/{resume_id}", headers=headers)
    assert get_resp.status_code == 200

    delete_resp = client.delete(f"/api/resumes/{resume_id}", headers=headers)
    assert delete_resp.status_code == 204

    get_after_delete = client.get(f"/api/resumes/{resume_id}", headers=headers)
    assert get_after_delete.status_code == 404


def test_cannot_access_another_users_resume(client):
    headers_a = _signup_and_login(client, email="ownera@example.com")
    headers_b = _signup_and_login(client, email="ownerb@example.com")

    docx_bytes = _make_docx_bytes()
    upload_resp = client.post(
        "/api/resumes/upload",
        headers=headers_a,
        files={"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    resume_id = upload_resp.json()["id"]

    resp = client.get(f"/api/resumes/{resume_id}", headers=headers_b)
    assert resp.status_code == 404


def test_analyze_resume_without_job_description_uses_fallback_engine(client):
    headers = _signup_and_login(client, email="analyze1@example.com")
    docx_bytes = _make_docx_bytes()
    upload_resp = client.post(
        "/api/resumes/upload",
        headers=headers,
        files={"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    resume_id = upload_resp.json()["id"]

    resp = client.post(f"/api/resumes/{resume_id}/analyze", headers=headers, json={})
    assert resp.status_code == 201
    data = resp.json()
    # No GEMINI_API_KEY configured in CI/test env -> must be the heuristic engine
    assert data["engine"] == "heuristic"
    assert 0 <= data["ats_score"] <= 100


def test_analyze_resume_with_keyword_matching_job_description(client):
    headers = _signup_and_login(client, email="analyze2@example.com")
    docx_bytes = _make_docx_bytes()
    upload_resp = client.post(
        "/api/resumes/upload",
        headers=headers,
        files={"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    resume_id = upload_resp.json()["id"]

    # Job description with clear taxonomy hits so this exercises the plain
    # keyword-overlap path, not the semantic-embedding fallback (keeps this
    # test fast and independent of any model download).
    resp = client.post(
        f"/api/resumes/{resume_id}/analyze",
        headers=headers,
        json={"job_description": "We need a Python developer with React and Docker experience."},
    )
    assert resp.status_code == 201
    data = resp.json()
    assert "python" in data["matched_skills"]
    assert "react" in data["matched_skills"]


def test_tailor_without_gemini_key_returns_503(client):
    headers = _signup_and_login(client, email="tailor1@example.com")
    docx_bytes = _make_docx_bytes()
    upload_resp = client.post(
        "/api/resumes/upload",
        headers=headers,
        files={"file": ("resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    resume_id = upload_resp.json()["id"]

    resp = client.post(
        f"/api/resumes/{resume_id}/tailor",
        headers=headers,
        json={"job_description": "Backend role requiring Python."},
    )
    # No GEMINI_API_KEY configured -> the endpoint must fail clearly, not silently
    assert resp.status_code == 503
